"""Excel 输出：结构化目录、映射关系表、盘点台账、质检结果。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..core.db import loads

MAPPING_COLUMNS = [
    ("档号", "dh"), ("源相对路径", "src_rel"), ("源文件名", "src_name"), ("是否目录", "is_dir"),
    ("文件数", "file_count"), ("类别代码", "category"), ("登记角色", "role"), ("源题名", "title"),
    ("题名来源", "title_source"), ("时期", "period"), ("项目UUID", "project_uuid"),
    ("项目层级", "project_level"), ("匹配状态", "match_status"), ("阶段", "stage"), ("子目录", "subdir"),
    ("文件类型", "file_type"), ("期间标识", "period_id"), ("版本", "version"),
    ("治理后相对路径", "target_rel"), ("治理后文件名", "target_name"), ("处理方式", "method"),
    ("去向状态", "status"), ("判定依据", "evidence"), ("置信度", "confidence"), ("规则版本", "rule_version"),
    ("密级", "secrecy"), ("涉密标记", "secret_flag"), ("真实性标记", "authenticity"),
    ("人工确认人", "confirmer"), ("复核人", "reviewer"), ("处理时间", "processed_at"), ("备注", "note"),
]

_HEAD_FONT = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")


def _write_sheet(ws, headers: list[str], rows: list[list], widths: dict | None = None) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.font = _HEAD_FONT
        c.fill = _HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        w = (widths or {}).get(h) or min(max(len(str(h)) * 2 + 2, 10), 40)
        ws.column_dimensions[get_column_letter(i)].width = w
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def _text_cells(ws, col_names: list[str], headers: list[str]) -> None:
    """档号等标识按文本格式保存，保留前导零。"""
    for name in col_names:
        if name in headers:
            idx = headers.index(name) + 1
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                for c in row:
                    c.number_format = "@"
                    if c.value is not None:
                        c.value = str(c.value)


def write_mapping(rows: list[dict], path: Path, rule_version: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "映射关系表"
    headers = [h for h, _ in MAPPING_COLUMNS]
    data = [[r.get(k) for _, k in MAPPING_COLUMNS] for r in rows]
    _write_sheet(ws, headers, data, {"源相对路径": 40, "治理后相对路径": 60, "源题名": 40, "判定依据": 40,
                                     "备注": 40, "治理后文件名": 50})
    _text_cells(ws, ["档号", "项目UUID"], headers)
    info = wb.create_sheet("说明")
    info.append(["规则版本", rule_version])
    info.append(["记录数", len(rows)])
    info.append(["说明", "每个被处理的源文件一行，包括未落库文件；去向状态分已落库、关联保存、原始区保留、待处理。"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def resolve_value(source: str, ctx: dict):
    prefix, _, key = source.partition(".")
    obj = ctx.get(prefix) or {}
    v = obj.get(key) if isinstance(obj, dict) else None
    return v if v not in (None, "") else None


def write_catalog(records: list[dict], header_cfg: dict, path: Path) -> None:
    """records 每项为 {archive, map, doc, case, project} 上下文。"""
    cols = header_cfg["columns"]
    missing = header_cfg.get("missing_mark", "【缺失】")
    wb = Workbook()
    ws = wb.active
    ws.title = header_cfg.get("sheet_name", "结构化目录")
    headers = [c["name"] for c in cols]
    data = []
    src_rows = []
    for rec in records:
        row, srow = [], []
        for c in cols:
            val, used = None, None
            for s in c["sources"]:
                val = resolve_value(s, rec)
                if val is not None:
                    used = s
                    break
            row.append(val if val is not None else missing)
            srow.append(used or "缺失")
        data.append(row)
        src_rows.append(srow)
    _write_sheet(ws, headers, data, {"文件题名": 40, "资料题名": 40, "落库文件名": 50, "落库相对路径": 60,
                                     "项目名称": 36})
    _text_cells(ws, ["档号", "项目编号", "项目UUID"], headers)
    ws2 = wb.create_sheet("取值来源")
    _write_sheet(ws2, headers, src_rows)
    ws3 = wb.create_sheet("表头配置")
    _write_sheet(ws3, ["列名", "取值来源优先级"], [[c["name"], " > ".join(c["sources"])] for c in cols])
    ws3.append([])
    ws3.append(["表头版本", header_cfg.get("version")])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_inventory(rows: list, path: Path) -> None:
    headers = ["档号", "时期", "时期代码", "识别得分", "目录入口", "实扫文件数", "实扫文件夹数", "容量MB",
               "类别分布", "对账状态", "对账差异", "处理状态", "处理说明", "规则版本", "数据路径"]
    data = []
    for r in rows:
        cc = loads(r["category_counts"], {})
        data.append([r["dh"], r["period"], r["period_code"], r["period_score"],
                     "；".join(loads(r["catalog_entries"], [])[:6]), r["file_count"], r["dir_count"],
                     round((r["total_bytes"] or 0) / 1024 ** 2, 2),
                     " ".join(f"{k}{v}" for k, v in (cc.get("成果") or {}).items()),
                     r["recon_status"], "；".join(loads(r["recon_detail"], [])), r["proc_status"],
                     r["proc_message"], r["proc_rule_version"], r["root"]])
    wb = Workbook()
    ws = wb.active
    ws.title = "盘点台账"
    _write_sheet(ws, headers, data, {"对账差异": 50, "数据路径": 50, "目录入口": 40, "处理说明": 30})
    _text_cells(ws, ["档号"], headers)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_qc(summary: list[dict], failures: list[dict], path: Path, meta: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "校验汇总"
    _write_sheet(ws, ["校验项", "应检数", "通过数", "未通过数", "结论"],
                 [[s["check"], s["total"], s["total"] - s["failed"], s["failed"],
                   "通过" if s["failed"] == 0 else "未通过"] for s in summary], {"校验项": 22})
    ws.append([])
    for k, v in meta.items():
        ws.append([k, v])
    ws2 = wb.create_sheet("未通过明细")
    _write_sheet(ws2, ["校验项", "档号", "对象", "说明"],
                 [[f["check"], f.get("dh"), f["target"], f["message"]] for f in failures],
                 {"对象": 70, "说明": 60})
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
